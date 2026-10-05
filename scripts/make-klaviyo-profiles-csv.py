#!/usr/bin/env python3
"""
Builds a Klaviyo profile-import CSV (Lists & Segments > Create list > Import) from the test-order workbook.
One row per customer. Emails are mindx.digitalmarketing+name@gmail.com plus-addresses, so any Klaviyo email
lands in your own inbox.

    python scripts/make-klaviyo-profiles-csv.py [scripts/out/rs1_orders_matrixify_100.xlsx]
    -> scripts/out/klaviyo_profiles_100.csv
"""
import csv
import sys
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "scripts/out/rs1_orders_matrixify_100.xlsx"
ws = load_workbook(src).active
h = [c.value for c in ws[1]]
ix = h.index

people = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[ix("Command")] != "NEW":
        continue
    email = r[ix("Customer: Email")]
    p = people.setdefault(email, {"orders": 0, "channels": set(), "first": r[ix("Customer: First Name")], "last": r[ix("Customer: Last Name")],
                                  "phone": r[ix("Customer: Phone")], "city": r[ix("Shipping: City")], "region": r[ix("Shipping: Province Code")],
                                  "zip": r[ix("Shipping: Zip")]})
    p["orders"] += 1
    p["channels"].add(r[ix("Source")])

out = root / "scripts/out" / f"klaviyo_profiles_{src.stem.rsplit('_', 1)[-1]}.csv"
with open(out, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["Email", "First Name", "Last Name", "Phone Number", "City", "Region", "Zip", "Country", "Test Orders", "Sales Channels", "Source", "Tags"])
    for email, p in people.items():
        w.writerow([email, p["first"], p["last"], p["phone"], p["city"], p["region"], p["zip"], "India", p["orders"],
                    ", ".join(sorted(p["channels"])), "RS1 test data", "TEST-DATA"])
print(f"Wrote {len(people)} profiles to {out}")
