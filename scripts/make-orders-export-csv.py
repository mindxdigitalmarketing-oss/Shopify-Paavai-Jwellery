#!/usr/bin/env python3
"""
Reads the Matrixify test-order workbook and writes a CSV in Shopify's own orders-export format
(one header row per order), which integrations/gorgias_sample_tickets.py expects.

    python scripts/make-orders-export-csv.py [scripts/out/rs1_orders_matrixify_100.xlsx]
    -> scripts/out/rs1_orders_export_100.csv
"""
import csv
import sys
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "scripts/out/rs1_orders_matrixify_100.xlsx"
ws = load_workbook(src).active
h = [c.value for c in ws[1]]
ix = h.index

orders = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    name = r[0]
    o = orders.setdefault(name, {"items": [], "fulfilled": False, "paid_amount": None})
    if r[ix("Command")] == "NEW":
        o.update(
            email=r[ix("Customer: Email")], created=r[ix("Created At")], financial=r[ix("Financial Status")], tags=r[ix("Tags")], source=r[ix("Source")],
            name=f"{r[ix('Shipping: First Name')]} {r[ix('Shipping: Last Name')]}", city=r[ix("Shipping: City")], zip=r[ix("Shipping: Zip")],
            phone=r[ix("Shipping: Phone")], ship=float(r[ix("Shipping Line: Price")] or 0), currency=r[ix("Currency")],
        )
    if r[ix("Line: Type")] == "Line Item":
        o["items"].append((r[ix("Line: Title")], int(r[ix("Line: Quantity")]), float(r[ix("Line: Price")])))
    if r[ix("Fulfillment: Status")]:
        o["fulfilled"] = True
        o["fulfilled_at"] = r[ix("Fulfillment: Processed At")]

out = root / "scripts/out" / (src.stem.replace("matrixify", "export") + ".csv")
cols = ["Name", "Email", "Financial Status", "Fulfillment Status", "Currency", "Total", "Created at", "Shipping Name", "Shipping City", "Shipping Zip",
        "Billing Name", "Phone", "Lineitem name", "Lineitem quantity", "Lineitem price", "Payment Method", "Tags", "Source", "Refunded Amount"]
with open(out, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, cols)
    w.writeheader()
    for name, o in orders.items():
        total = sum(q * p for _, q, p in o["items"]) + o["ship"]
        title, qty, price = o["items"][0]
        w.writerow({
            "Name": name, "Email": o["email"], "Financial Status": o["financial"], "Fulfillment Status": "fulfilled" if o["fulfilled"] else "unfulfilled",
            "Currency": o["currency"], "Total": f"{total:.2f}", "Created at": o["created"].replace("T", " ").replace(".000Z", " +0000"),
            "Shipping Name": o["name"], "Shipping City": o["city"], "Shipping Zip": o["zip"], "Billing Name": o["name"], "Phone": o["phone"],
            "Lineitem name": title, "Lineitem quantity": qty, "Lineitem price": f"{price:.2f}",
            "Payment Method": "manual" if o["financial"] == "paid" else "Cash on Delivery (COD)", "Tags": o["tags"], "Source": o["source"], "Refunded Amount": "0.00",
        })
        for title, qty, price in o["items"][1:]:
            w.writerow({"Name": name, "Lineitem name": title, "Lineitem quantity": qty, "Lineitem price": f"{price:.2f}"})
print(f"Wrote {len(orders)} orders to {out}")
